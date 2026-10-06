"""Backtest reports (one or several timeframes): Telegram summary, Excel workbook (Arabic, RTL) and CSV."""
from __future__ import annotations

import csv
import html
import io
from datetime import datetime, timezone
from typing import Any

from app.backtest.engine import BtResult, Trade

Results = list[tuple[str, BtResult]]      # [(timeframe, result)] in timeframe order

TRADE_HEADERS = ["#", "الزوج", "وقت الدخول", "TF", "اتجاه", "اللوت", "سعر الدخول", "الهدف (TP)", "الوقف (SL)", "وقت الخروج", "سعر الخروج",
                 "سبب الخروج", "النتيجة", "النقاط", "الربح الإجمالي ($)", "العمولة ($)", "صافي الربح ($)", "رصيد تراكمي ($)",
                 "أقصى تراجع بالصفقة (نقاط)", "أقصى ربح عائم (نقاط)", "عدد الشموع", "تعليق الإشارة"]
REASON_AR = {"TP": "TP ✅ هدف", "SL": "SL ❌ وقف", "SIGNAL_CLOSE": "إغلاق بإشارة", "DAILY_LOSS_LIMIT": "🛑 حد خسارة يومي",
             "DAILY_PROFIT_TARGET": "✅ هدف ربح يومي", "END_OF_TEST": "نهاية الاختبار"}


def _fmt(ts: int, tz_ms: int) -> str:
    return datetime.fromtimestamp((ts + tz_ms) / 1000, timezone.utc).strftime("%Y-%m-%d %H:%M")


def _result(t: Trade) -> str:
    return "WIN ✅" if t.net > 0 else ("LOSS ❌" if t.net < 0 else "BE ➖")


def _num(v: float | None, d: int = 2) -> str:
    return "—" if v is None else f"{v:,.{d}f}"


def trade_row(t: Trade, tf: str, tz_ms: int) -> list[Any]:
    return [t.n, t.symbol, _fmt(t.entry_ts, tz_ms), tf, "BUY 📈" if t.side == "BUY" else "SELL 📉", t.lots, t.entry_price, t.tp, t.sl,
            _fmt(t.exit_ts, tz_ms), t.exit_price, REASON_AR.get(t.reason, t.reason), _result(t), round(t.pips, 1), round(t.gross, 2),
            round(t.commission, 2), round(t.net, 2), round(t.balance_after, 2), round(t.mae_pips, 1), round(t.mfe_pips, 1), t.bars, t.comment]


def _price_fmt(decimals: int) -> str:
    return "0" if decimals <= 0 else "0." + "0" * decimals


def metric_rows(r: BtResult) -> list[tuple[str, Any]]:
    s, ccy = r.stats, r.cfg.deposit_ccy
    pf = None if s["profit_factor"] is None else round(s["profit_factor"], 2)
    return [
        (f"الرصيد الابتدائي ({ccy})", round(s["initial_balance"], 2)), (f"الرصيد النهائي ({ccy})", round(s["final_balance"], 2)),
        (f"صافي الربح ({ccy})", round(s["net_profit"], 2)), ("العائد %", round(s["return_pct"], 2)),
        (f"إجمالي الربح ({ccy})", round(s["gross_profit"], 2)), (f"إجمالي الخسارة ({ccy})", round(s["gross_loss"], 2)),
        ("معامل الربح (Profit Factor)", pf), ("عدد الصفقات", s["total_trades"]), ("رابحة", s["wins"]), ("خاسرة", s["losses"]),
        ("متعادلة", s["breakeven"]), ("نسبة النجاح %", round(s["win_rate"], 1)), ("متوسط الربح", round(s["avg_win"], 2)),
        ("متوسط الخسارة", round(s["avg_loss"], 2)), ("نسبة الربح/الخسارة (Payoff)", None if s["payoff_ratio"] is None else round(s["payoff_ratio"], 2)),
        ("العائد المتوقع للصفقة", round(s["expected_payoff"], 2)), ("أكبر ربح", round(s["largest_win"], 2)), ("أكبر خسارة", round(s["largest_loss"], 2)),
        ("أطول سلسلة ربح (عدد)", s["max_consec_wins"]), ("أطول سلسلة خسارة (عدد)", s["max_consec_losses"]),
        ("إجمالي النقاط", round(s["total_pips"], 1)), ("متوسط النقاط/صفقة", round(s["avg_pips"], 1)),
        ("أقصى تراجع (قيمة)", round(s["max_drawdown"], 2)), ("أقصى تراجع %", round(s["max_drawdown_pct"], 2)),
        ("عامل التعافي", None if s["recovery_factor"] is None else round(s["recovery_factor"], 2)),
        ("نسبة شارب (يومية، سنوية)", None if s["sharpe"] is None else round(s["sharpe"], 2)),
        ("صفقات شراء", s["buys"]), ("ربح الشراء", round(s["buy_net"], 2)), ("صفقات بيع", s["sells"]), ("ربح البيع", round(s["sell_net"], 2)),
        ("متوسط مدة الصفقة (ساعات)", round(s["avg_hold_hours"], 1)), ("إشارات صدرت", r.signals_total), ("إشارات متجاهلة", len(r.skipped)),
    ]


def info_rows(meta: dict[str, Any]) -> list[tuple[str, Any]]:
    return [("الاستراتيجية", f"{meta['strategy']} v{meta['version']}"), ("الحساب (مصدر البيانات)", meta["account"]),
            ("الرموز", ", ".join(meta["symbols"])), ("الفريمات", ", ".join(meta["timeframes"])),
            ("الفترة", f"{meta['date_from']} → {meta['date_to']} ({meta['tz_label']})"), ("عدد الشموع المعالجة", meta["bars"]),
            ("مدة التشغيل (ثانية)", round(meta.get("seconds", 0), 1))]


def settings_rows(results: Results, meta: dict[str, Any]) -> list[tuple[str, Any]]:
    o = meta.get("options", {})
    rows: list[tuple[str, Any]] = [
        ("الرصيد الابتدائي", o.get("initial_balance")), ("عملة الحساب", results[0][1].cfg.deposit_ccy),
        ("حجم اللوت", o.get("lot_size") or "حسب الاستراتيجية (احتياطي 0.10)"),
        ("وقف الخسارة (نقاط)", o.get("sl_points") or "حسب الاستراتيجية"), ("الهدف / Take Profit (نقاط)", o.get("tp_points") or "حسب الاستراتيجية"),
        ("السبريد", o.get("spread")), ("فارق التوقيت (ساعة)", o.get("tz_offset"))]
    for sym, sp in meta.get("spreads", {}).items():
        rows.append((f"السبريد المستخدم {sym} (نقاط)", sp))
    rows += [("", ""), ("تعريف النقطة (Point)", "")]
    for sym, d in meta.get("points", {}).items():
        rows.append((sym, f"1 نقطة = {d['point']} · المنازل المعروضة {d['decimals']} · قيمة النقطة للوت الواحد ≈ {d['value']:.4g} {d['ccy']}"))
    rows += [("", ""), ("إعدادات الاستراتيجية المستخدمة", "")] + [(k, v) for k, v in meta.get("params", {}).items()]
    rows += [("", ""), ("افتراضات المحاكاة", ""),
             ("تنفيذ الإشارة", "عند افتتاح الشمعة التالية بعد إغلاق شمعة الإشارة"),
             ("الأسعار", "Bid من الشارت؛ الشراء بسعر Ask = Open + السبريد، البيع بسعر Bid"),
             ("SL و TP في نفس الشمعة", "يُحتسب الوقف أولًا (تقدير متحفظ)"),
             ("الفجوات", "الوقف يُنفذ بسعر الافتتاح؛ الهدف بسعر الافتتاح إن كان أفضل"),
             ("غير محاكى", "السواب، الانزلاق، الهامش/Stop-out، الأوامر المعلقة، قيود محرك المخاطر الحي"),
             ("دقة البيانات", "شموع OHLC من وسيطك عبر cTrader؛ كلما صغر الفريم زادت الدقة")]
    for tf, r in results:
        for w in r.warnings:
            rows.append((f"تحذير [{tf}]", w))
    return rows


# ---- Telegram summary ------------------------------------------------------------------------------------------------
def build_summary_text(results: Results, meta: dict[str, Any]) -> str:
    e = html.escape
    head = [f"🧪 <b>Backtest — {e(meta['strategy'])} v{e(meta['version'])}</b>", f"{e(', '.join(meta['symbols']))}",
            f"{e(meta['date_from'])} → {e(meta['date_to'])} ({e(meta['tz_label'])}) · {meta['bars']:,} bars · {meta.get('seconds', 0):.0f}s", ""]
    if len(results) == 1:
        tf, r = results[0]
        s = r.stats
        pf = "∞" if s["profit_factor"] is None and s["gross_profit"] > 0 else _num(s["profit_factor"])
        head[1] += f" · <b>{e(tf)}</b>"
        lines = head + [
            f"💰 الرصيد: {_num(s['initial_balance'])} → <b>{_num(s['final_balance'])}</b> {e(r.cfg.deposit_ccy)}",
            f"{'🟢' if s['net_profit'] >= 0 else '🔴'} صافي الربح: <b>{s['net_profit']:+,.2f}</b> ({s['return_pct']:+.2f}%)",
            f"📊 الصفقات: <b>{s['total_trades']}</b> (✅ {s['wins']} / ❌ {s['losses']}) · النجاح <b>{s['win_rate']:.1f}%</b>",
            f"⚖️ Profit Factor: <b>{pf}</b> · Payoff: {_num(s['payoff_ratio'])} · Expectancy: {s['expected_payoff']:+.2f}",
            f"📉 أقصى تراجع: <b>{_num(s['max_drawdown'])}</b> ({s['max_drawdown_pct']:.2f}%) · Recovery: {_num(s['recovery_factor'])}",
            f"🏆 أكبر ربح {s['largest_win']:+.2f} · أكبر خسارة {s['largest_loss']:+.2f}",
            f"🔁 أطول ربح متتالٍ {s['max_consec_wins']} · أطول خسارة متتالية {s['max_consec_losses']}",
            f"📐 النقاط: {s['total_pips']:+.1f} (متوسط {s['avg_pips']:+.1f}) · شارب {_num(s['sharpe'])}"]
        if r.per_symbol:
            lines += ["", "<b>حسب الرمز</b>"] + [f"• {e(p['symbol'])}: {p['trades']} صفقة · {p['win_rate']:.0f}% · {p['net']:+,.2f}" for p in r.per_symbol[:12]]
    else:
        lines = head + ["<b>مقارنة الفريمات</b>  (صافي · صفقات · نجاح · PF · تراجع)"]
        best = max(results, key=lambda x: x[1].stats["net_profit"])
        for tf, r in results:
            s = r.stats
            pf = "∞" if s["profit_factor"] is None and s["gross_profit"] > 0 else _num(s["profit_factor"], 1)
            lines.append(f"{'🏆' if tf == best[0] else '•'} <b>{e(tf)}</b>: {s['net_profit']:+,.0f} · {s['total_trades']} · {s['win_rate']:.0f}% · {pf} · {s['max_drawdown_pct']:.1f}%")
        lines += ["", f"الأفضل صافيًا: <b>{e(best[0])}</b> ({best[1].stats['net_profit']:+,.2f} {e(best[1].cfg.deposit_ccy)})", "التفاصيل الكاملة لكل فريم داخل ملف Excel."]
    warns = [f"{tf}: {w}" for tf, r in results for w in r.warnings][:4]
    lines += [f"⚠️ {e(w)}" for w in warns]
    lines.append("\n<i>محاكاة على شموع OHLC؛ لا تضمن نتائج مستقبلية.</i>")
    return "\n".join(lines)[:3900]


# ---- CSV --------------------------------------------------------------------------------------------------------------
def build_csv(results: Results, meta: dict[str, Any]) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(TRADE_HEADERS)
    for tf, r in results:
        tz_ms = int(r.cfg.tz_offset_hours * 3_600_000)
        for t in r.trades:
            w.writerow(trade_row(t, tf, tz_ms))
    return ("\ufeff" + buf.getvalue()).encode("utf-8")        # BOM: Excel shows Arabic correctly


def _sample(points: list, limit: int = 2500) -> list:
    if len(points) <= limit:
        return points
    step = len(points) / limit
    return [points[int(i * step)] for i in range(limit)] + [points[-1]]


# ---- Excel ------------------------------------------------------------------------------------------------------------
def build_xlsx(results: Results, meta: dict[str, Any]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    multi = len(results) > 1
    tz_ms = int(results[0][1].cfg.tz_offset_hours * 3_600_000)
    base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    head_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="1F4E78")
    green, red = Font(name="Arial", size=10, color="0B7A32"), Font(name="Arial", size=10, color="B00020")
    thin = Side(style="thin", color="D0D0D0")
    border = Border(top=thin, bottom=thin, left=thin, right=thin)
    center = Alignment(horizontal="center", vertical="center")
    dec = meta.get("display_decimals", {})

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(title: str, headers: list[str], rows: list[list[Any]], widths: list[int] | None = None, color_col: int | None = None) -> Any:
        ws = wb.create_sheet(title)
        ws.sheet_view.rightToLeft = True
        ws.append(headers)
        for c in range(1, len(headers) + 1):
            cell = ws.cell(1, c)
            cell.font, cell.fill, cell.alignment, cell.border = head_font, head_fill, center, border
        for row in rows:
            ws.append(row)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                cell.font, cell.border, cell.alignment = base, border, center
            if color_col is not None:
                v = row[color_col - 1].value
                if isinstance(v, (int, float)):
                    row[color_col - 1].font = green if v > 0 else (red if v < 0 else base)
                elif v is not None:
                    row[color_col - 1].font = green if "WIN" in str(v) else (red if "LOSS" in str(v) else base)
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1] if widths else max(12, min(32, len(str(h)) + 4))
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        return ws

    # 1. summary: general info + metrics table (one column per timeframe)
    ws = wb.create_sheet("الملخص")
    ws.sheet_view.rightToLeft = True
    for k, v in info_rows(meta):
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font, ws.cell(ws.max_row, 2).font = bold, base
    ws.append([])
    hdr_row = ws.max_row + 1
    ws.append(["المؤشر"] + [tf for tf, _ in results])
    for c in range(1, len(results) + 2):
        ws.cell(hdr_row, c).font, ws.cell(hdr_row, c).fill, ws.cell(hdr_row, c).alignment = head_font, head_fill, center
    per_tf = [metric_rows(r) for _, r in results]
    for i, (label, _) in enumerate(per_tf[0]):
        ws.append([label] + [rows[i][1] for rows in per_tf])
        for c in range(1, len(results) + 2):
            ws.cell(ws.max_row, c).font = base
            ws.cell(ws.max_row, c).alignment = center if c > 1 else Alignment(horizontal="right")
            ws.cell(ws.max_row, c).border = border
    ws.column_dimensions["A"].width = 38
    for i in range(2, len(results) + 2):
        ws.column_dimensions[get_column_letter(i)].width = 22 if i == 2 else 16
    ws.freeze_panes = None

    # 2. trades (all timeframes; TF column) with price formats per symbol (gold without decimals)
    trade_rows = []
    for tf, r in results:
        trade_rows += [trade_row(t, tf, tz_ms) for t in r.trades]
    ws_t = sheet("الصفقات", TRADE_HEADERS, trade_rows, [6, 12, 17, 6, 11, 8, 12, 12, 12, 17, 12, 18, 11, 10, 14, 11, 14, 15, 14, 14, 9, 28], color_col=13)
    for rr in range(2, ws_t.max_row + 1):
        pf = _price_fmt(dec.get(ws_t.cell(rr, 2).value, 5))
        for c in (7, 8, 9, 11):
            ws_t.cell(rr, c).number_format = pf
        ws_t.cell(rr, 6).number_format = "0.00"
        ws_t.cell(rr, 14).number_format = "0.0"
        for c in (15, 16, 17, 18):
            ws_t.cell(rr, c).number_format = "#,##0.00"
        v = ws_t.cell(rr, 17).value
        if isinstance(v, (int, float)):
            ws_t.cell(rr, 17).font = green if v > 0 else (red if v < 0 else base)

    # 3-5. by symbol / daily / monthly (TF column)
    sheet("حسب الرمز", ["TF", "الرمز", "الصفقات", "رابحة", "خاسرة", "نسبة النجاح %", "صافي الربح ($)", "النقاط", "Profit Factor", "متوسط الصفقة ($)"],
          [[tf, p["symbol"], p["trades"], p["wins"], p["losses"], round(p["win_rate"], 1), round(p["net"], 2), round(p["pips"], 1),
            None if p["profit_factor"] is None else round(p["profit_factor"], 2), round(p["avg_net"], 2)] for tf, r in results for p in r.per_symbol], color_col=7)
    sheet("يومي", ["TF", "التاريخ", "الصفقات", "الرابحة", "صافي اليوم ($)", "الرصيد ($)"],
          [[tf, d["date"], d["trades"], d["wins"], round(d["net"], 2), round(d["balance"], 2)] for tf, r in results for d in r.daily], color_col=5)
    sheet("شهري", ["TF", "الشهر", "الصفقات", "الرابحة", "صافي الشهر ($)", "العائد %", "الرصيد ($)"],
          [[tf, m["month"], m["trades"], m["wins"], round(m["net"], 2), round(m["return_pct"], 2), round(m["balance"], 2)] for tf, r in results for m in r.monthly], color_col=5)
    stopped = [[tf, d["date"], d["reason"]] for tf, r in results for d in r.stopped_days]
    if stopped:
        sheet("أيام الإيقاف", ["TF", "التاريخ", "السبب (النتيجة)"], stopped, [8, 14, 120])
    sheet("إشارات متجاهلة", ["TF", "الوقت", "الرمز", "الاتجاه", "السبب"],
          [[tf, _fmt(s["ts"], tz_ms), s["symbol"], s["side"], s["reason"]] for tf, r in results for s in r.skipped[:5000]], [8, 18, 12, 10, 60])
    sheet("منحنى الرصيد", ["TF", "الوقت", "الرصيد ($)", "Equity ($)"],
          [[tf, _fmt(ts, tz_ms), round(b, 2), round(eq, 2)] for tf, r in results for ts, b, eq in _sample(r.equity)], [8, 18, 16, 16])

    ws_c = wb.create_sheet("الإعدادات")
    ws_c.sheet_view.rightToLeft = True
    ws_c.append(["البند", "القيمة"])
    for c in (1, 2):
        ws_c.cell(1, c).font, ws_c.cell(1, c).fill = head_font, head_fill
    for k, v in settings_rows(results, meta):
        ws_c.append([k, v if not isinstance(v, (dict, list)) else str(v)])
    for row in ws_c.iter_rows(min_row=2, max_row=ws_c.max_row):
        for cell in row:
            cell.font, cell.alignment = base, Alignment(wrap_text=True, vertical="top")
    ws_c.column_dimensions["A"].width, ws_c.column_dimensions["B"].width = 42, 90
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
