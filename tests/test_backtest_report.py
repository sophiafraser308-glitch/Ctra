import csv
import io
import os
import sys

import pytest

pytest.importorskip("openpyxl")
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))
from app.backtest.engine import BacktestEngine, BtConfig, SymbolSpec  # noqa: E402
from app.backtest.report import TRADE_HEADERS, build_csv, build_summary_text, build_xlsx  # noqa: E402

TF = 900_000


def make_result():
    cfg = BtConfig(tf_ms=TF, tz_offset_hours=3, daily_loss_limit=0)
    e = BacktestEngine(cfg, {"EURUSD": SymbolSpec("EURUSD", 0.0001, 100_000, 5, spread_pips=1.0)})
    sig = {"side": "BUY", "order_type": "MARKET", "stop_loss_pips": 10, "take_profit_pips": 20, "volume_lots": 1.0}
    for k in range(10):
        e.feed(2 * k * TF, "EURUSD", {"t": 2 * k * TF, "o": 1.1, "h": 1.1, "l": 1.1, "c": 1.1, "v": 1}, [sig])
        win = k % 2 == 0
        e.feed((2 * k + 1) * TF, "EURUSD", {"t": (2 * k + 1) * TF, "o": 1.1, "h": 1.1030 if win else 1.1, "l": 1.0999 if win else 1.098, "c": 1.1, "v": 1}, [])
    return e.finish()


META = {"strategy": "Stochastic Cross", "version": "1.0.0", "account": "Demo", "symbols": ["EURUSD"], "timeframes": ["M15", "H1"], "date_from": "2026-01-01",
        "date_to": "2026-01-31", "tz_label": "UTC+3", "bars": 20, "seconds": 1.2, "params": {"k_period": 49}, "spreads": {"EURUSD": 1.0}, "options": {"initial_balance": 10000, "lot_size": 0, "sl_points": 0, "tp_points": 0, "spread": "auto", "tz_offset": 3}, "points": {"EURUSD": {"point": 0.0001, "decimals": 5, "ccy": "USD", "value": 10.0}}, "display_decimals": {"EURUSD": 5}}


RESULTS = None


def results():
    global RESULTS
    if RESULTS is None:
        RESULTS = [("M15", make_result()), ("H1", make_result())]
    return RESULTS


def test_xlsx_multi_timeframe():
    from openpyxl import load_workbook
    rs = results()
    wb = load_workbook(io.BytesIO(build_xlsx(rs, META)))
    assert wb.sheetnames == ["الملخص", "الصفقات", "حسب الرمز", "يومي", "شهري", "إشارات متجاهلة", "منحنى الرصيد", "الإعدادات"]
    ws = wb["الصفقات"]
    assert [c.value for c in ws[1]] == TRADE_HEADERS
    assert ws.max_row == 1 + sum(len(r.trades) for _, r in rs)
    assert {ws.cell(i, 4).value for i in range(2, ws.max_row + 1)} == {"M15", "H1"}
    summ = wb["الملخص"]
    header = [c.value for c in summ[9]]                       # info rows (7) + blank + header
    assert header[:3] == ["المؤشر", "M15", "H1"]


def test_gold_prices_shown_without_decimals():
    from openpyxl import load_workbook
    from app.backtest.engine import BacktestEngine, BtConfig, SymbolSpec
    e = BacktestEngine(BtConfig(tf_ms=TF, tz_offset_hours=0), {"XAUUSD": SymbolSpec("XAUUSD", 0.1, 100, 2, spread_pips=0.0, display_decimals=0)})
    sig = {"side": "BUY", "order_type": "MARKET", "stop_loss_pips": 20, "take_profit_pips": 20, "volume_lots": 1.0}
    e.feed(0, "XAUUSD", {"t": 0, "o": 4150.0, "h": 4150.0, "l": 4150.0, "c": 4150.0, "v": 1}, [sig])
    e.feed(TF, "XAUUSD", {"t": TF, "o": 4150.0, "h": 4152.5, "l": 4149.9, "c": 4152.0, "v": 1}, [])
    r = e.finish()
    meta = dict(META, symbols=["XAUUSD"], display_decimals={"XAUUSD": 0})
    ws = load_workbook(io.BytesIO(build_xlsx([("M15", r)], meta)))["الصفقات"]
    assert ws["G2"].number_format == "0" and ws["K2"].number_format == "0" and ws["N2"].value == 20.0


def test_csv_and_summary():
    rs = results()
    rows = list(csv.reader(io.StringIO(build_csv(rs, META).decode("utf-8-sig"))))
    assert rows[0] == TRADE_HEADERS and len(rows) == 1 + sum(len(r.trades) for _, r in rs)
    text = build_summary_text(rs, META)
    assert "مقارنة الفريمات" in text and "M15" in text and "H1" in text and len(text) < 4000
    single = build_summary_text(rs[:1], dict(META, timeframes=["M15"]))
    assert "حسب الرمز" in single
